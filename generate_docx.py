"""Generate a step-by-step DOCX guide: AI Use Case Intake to Go-To-Market.
Redesigned: multi-color, concise, role-per-task tables, scannable layout."""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import datetime

doc = Document()

# -- Page margins --
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

# -- Style setup --
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15

# Heading styles
h1 = doc.styles['Heading 1']
h1.font.name = 'Calibri'
h1.font.size = Pt(22)
h1.font.bold = True
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(6)

h2 = doc.styles['Heading 2']
h2.font.name = 'Calibri'
h2.font.size = Pt(14)
h2.font.bold = True
h2.paragraph_format.space_before = Pt(12)
h2.paragraph_format.space_after = Pt(4)

h3 = doc.styles['Heading 3']
h3.font.name = 'Calibri'
h3.font.size = Pt(11)
h3.font.bold = True
h3.paragraph_format.space_before = Pt(8)
h3.paragraph_format.space_after = Pt(3)

# Step colors (R, G, B)
COLORS = {
    1: RGBColor(0x25, 0x63, 0xEB),  # Blue
    2: RGBColor(0x10, 0xB9, 0x81),  # Emerald
    3: RGBColor(0xF5, 0x9E, 0x0B),  # Amber
    4: RGBColor(0xF4, 0x3F, 0x5E),  # Rose
    5: RGBColor(0x8B, 0x5C, 0xF6),  # Violet
    6: RGBColor(0x06, 0xB6, 0xD4),  # Cyan
    7: RGBColor(0xF9, 0x73, 0x16),  # Orange
    8: RGBColor(0x14, 0xB8, 0xA6),  # Teal
}

LIGHT_COLORS = {
    1: "D6E4FF",
    2: "D1FAE5",
    3: "FEF3C7",
    4: "FFE4E6",
    5: "EDE9FE",
    6: "CFFAFE",
    7: "FFEDD5",
    8: "CCFBF1",
}

GRAY = RGBColor(0x6B, 0x72, 0x80)


def set_cell_shading(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_task_table(doc, tasks, step_num):
    table = doc.add_table(rows=1 + len(tasks), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row in table.rows:
        row.cells[0].width = Cm(12)
        row.cells[1].width = Cm(4.5)

    # Header row
    hdr = table.rows[0]
    for i, text in enumerate(["Task", "Owner"]):
        cell = hdr.cells[i]
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = 'Calibri'
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        hex_c = str(COLORS[step_num])
        set_cell_shading(cell, hex_c)

    # Data rows
    for idx, (task, owner) in enumerate(tasks):
        row = table.rows[idx + 1]
        cell_t = row.cells[0]
        p = cell_t.paragraphs[0]
        parts = task.split("**")
        for j, part in enumerate(parts):
            run = p.add_run(part)
            run.font.size = Pt(10)
            run.font.name = 'Calibri'
            if j % 2 == 1:
                run.bold = True

        cell_o = row.cells[1]
        p2 = cell_o.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p2.add_run(owner)
        run.font.size = Pt(9)
        run.font.color.rgb = COLORS[step_num]
        run.font.name = 'Calibri'
        run.bold = True

        if idx % 2 == 0:
            set_cell_shading(cell_t, LIGHT_COLORS[step_num])
            set_cell_shading(cell_o, LIGHT_COLORS[step_num])

    doc.add_paragraph()


def add_step_heading(doc, step_num, title, phase):
    p = doc.add_heading(f"Step {step_num}: {title}", level=1)
    for run in p.runs:
        run.font.color.rgb = COLORS[step_num]

    p2 = doc.add_paragraph()
    run = p2.add_run(f"\u25CF  {phase.upper()}")
    run.font.size = Pt(9)
    run.font.color.rgb = COLORS[step_num]
    run.bold = True
    run.font.name = 'Calibri'

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(6)
    run = p3.add_run("\u2500" * 65)
    run.font.color.rgb = COLORS[step_num]
    run.font.size = Pt(6)


def add_gate(doc, title, subtitle, is_governance=False):
    color = RGBColor(0xDC, 0x26, 0x26) if is_governance else RGBColor(0xF5, 0x9E, 0x0B)
    light = "FEE2E2" if is_governance else "FEF3C7"

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, light)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\u26A0  {title}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = color
    run.font.name = 'Calibri'

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run(subtitle)
    run.font.size = Pt(9.5)
    run.font.color.rgb = GRAY
    run.font.name = 'Calibri'

    doc.add_paragraph()


# ============================================================
# COVER PAGE
# ============================================================
for _ in range(5):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("AI Use Case Intake")
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
run.bold = True
run.font.name = 'Calibri'

title2 = doc.add_paragraph()
title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title2.add_run("to Go-To-Market")
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
run.bold = True
run.font.name = 'Calibri'

doc.add_paragraph()

# Color bar
bar = doc.add_paragraph()
bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
for c in COLORS.values():
    run = bar.add_run("\u2588\u2588\u2588\u2588 ")
    run.font.color.rgb = c
    run.font.size = Pt(14)

doc.add_paragraph()

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Step-by-Step Guide")
run.font.size = Pt(18)
run.font.color.rgb = GRAY
run.font.name = 'Calibri'

doc.add_paragraph()

desc = doc.add_paragraph()
desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = desc.add_run(
    "8 steps from ideation to launch\n"
    "Who does what at every stage"
)
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
run.font.name = 'Calibri'

for _ in range(4):
    doc.add_paragraph()

footer_p = doc.add_paragraph()
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run("AI Power Team")
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
run.bold = True

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = date_p.add_run(datetime.date.today().strftime("%B %Y"))
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

doc.add_page_break()

# ============================================================
# TABLE OF CONTENTS
# ============================================================
toc_heading = doc.add_heading("Table of Contents", level=1)
for run in toc_heading.runs:
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

doc.add_paragraph()

toc_items = [
    (1, "Identify & Generate AI Use Cases", "Ideation"),
    (2, "Submit Use Case Intake", "Intake"),
    (3, "AI Enablement Team Review", "Assessment"),
    (0, "Release Gate Decision", ""),
    (4, "Project Setup & Resource Allocation", "Setup"),
    (5, "Define the MVP & Build", "Build"),
    (0, "Governance Gate", ""),
    (6, "GTM Ready: Package the Solution", "GTM Ready"),
    (7, "GTM Push: Showcase & Sales Activation", "Go-To-Market"),
    (8, "Track Results & Continuous Improvement", "Measure"),
]

for num, title_text, phase in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    if num:
        run = p.add_run(f"  {num}.  ")
        run.font.size = Pt(13)
        run.font.color.rgb = COLORS[num]
        run.bold = True
        run.font.name = 'Calibri'
        run = p.add_run(title_text)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        run.font.name = 'Calibri'
        if phase:
            run = p.add_run(f"   {phase}")
            run.font.size = Pt(9)
            run.font.color.rgb = COLORS[num]
            run.font.name = 'Calibri'
    else:
        run = p.add_run(f"        \u26A0  {title_text}")
        run.font.size = Pt(10)
        run.font.color.rgb = GRAY
        run.italic = True
        run.font.name = 'Calibri'

doc.add_page_break()

# ============================================================
# STEP 1
# ============================================================
add_step_heading(doc, 1, "Identify & Generate AI Use Cases", "Ideation")

add_task_table(doc, [
    ("**Talk to your Practice Lead** about client pain points and market gaps", "Practice Lead"),
    ("Brainstorm from in-flight projects, client conversations & market demand", "AI Champion"),
    ("Check **Demo Catalog** (79+ solutions, 10 domains) to avoid duplication. Interested in a solution? **Request a demo directly with the Product Owner**", "AI Champion"),
], 1)

doc.add_page_break()

# ============================================================
# STEP 2
# ============================================================
add_step_heading(doc, 2, "Submit Use Case Intake", "Intake")

add_task_table(doc, [
    ("Complete **AI Use-Case Intake Form** (one row per use case)", "AI Champion"),
    ("Answer: **What is the pain/problem?**", "AI Champion"),
    ("Answer: **Who feels the pain?** (user, buyer, team)", "AI Champion"),
    ("Answer: **Why does it matter?** (cost, speed, risk, growth)", "AI Champion"),
    ("Answer: **Why is it unsolved today?** (blockers, workarounds)", "AI Champion"),
    ("Provide: title, practice, persona, business value, revenue potential, target accounts", "AI Champion"),
    ("Assign priority: **High** / **Medium** / **Low**", "AI Champion + Practice Lead"),
], 2)

doc.add_page_break()

# ============================================================
# STEP 3
# ============================================================
add_step_heading(doc, 3, "AI Enablement Team Review", "Assessment")

add_task_table(doc, [
    ("Determine scope: **POC / MVP / Full Build** based on business value & tech depth", "AI Enablement Team"),
    ("Assess **strategic alignment** and data maturity", "AI Enablement Team"),
    ("**Confirm resource availability** across AI Forward Engineers", "AI Enablement Team"),
    ("Prioritize **top 1\u20132 use cases** per champion for development", "AI Enablement Team"),
], 3)

add_gate(doc, "RELEASE GATE", "Approved to proceed?  Yes \u2192 continue to build   |   No \u2192 return to ideation")

doc.add_page_break()

# ============================================================
# STEP 4
# ============================================================
add_step_heading(doc, 4, "Project Setup & Resource Allocation", "Setup")

add_task_table(doc, [
    ("**Asset Register** \u2014 Model the AI solution into asset register (MP and above)", "Business Sponsor"),
    ("**Set Up Jira** \u2014 Create project, epics, stories & sprint backlog", "AI Enablement Team"),
    ("**Allocate Resources** \u2014 Assign AI Forward Engineers via Retain", "AI Enablement Team"),
    ("**Set up healthy check-in cadence** \u2014 demo \u2192 feedback \u2192 refine cycles (keep AI Enablement in the loop)", "AI Champion + AI Fwd Engineer"),
], 4)

doc.add_page_break()

# ============================================================
# STEP 5
# ============================================================
add_step_heading(doc, 5, "Define the MVP & Build", "Build")

add_task_table(doc, [
    ("**Define MVP scope** \u2014 features, demo expectations & success criteria", "AI Fwd Engineer + AI Champion"),
    ("Develop through **sprint cycles** with iterative testing & user feedback", "AI Forward Engineer"),
    ("Provide **sample data & real-life use case scenarios** to guide the build", "AI Champion"),
    ("Build **demo environment** based on provided scenarios & data", "AI Forward Engineer"),
    ("Track: Backlog \u2192 Build In-Progress \u2192 **MVP Demo Complete**", "AI Enablement Team"),
], 5)

add_gate(doc,
    "GOVERNANCE GATE",
    "AI Governance Evaluation  \u2022  Compliance Checks  \u2022  Data Security Review  \u2022  Model Risk Validation",
    is_governance=True)

doc.add_page_break()

# ============================================================
# STEP 6
# ============================================================
add_step_heading(doc, 6, "GTM Ready \u2014 Package the Solution", "GTM Ready")

add_task_table(doc, [
    ("Prepare **demo materials**: demo script, mock data & demo video recording", "AI Fwd Engineer + AI Champion"),
    ("Finalize **demo script** and conduct full run-throughs", "AI Fwd Engineer + AI Champion"),
    ('Prepare **"AI vs. Non-AI" talk track** (competitive differentiation)', "AI Champion + Sales"),
    ("Create **one-pager**, pricing template & engagement model", "Domain Leader + Sales"),
    ("Map **target accounts & buyer personas**; register in Demo Catalog", "Sales + AI Enablement Team"),
], 6)

# Deliverables callout
p = doc.add_paragraph()
run = p.add_run("Deliverables checklist: ")
run.bold = True
run.font.size = Pt(10)
run.font.color.rgb = COLORS[6]
run = p.add_run("Polished demo  |  Demo script  |  Mock data  |  Demo video  |  One-pager  |  Talk track  |  Pricing template  |  Target account list")
run.font.size = Pt(10)
run.font.color.rgb = GRAY

doc.add_page_break()

# ============================================================
# STEP 7
# ============================================================
add_step_heading(doc, 7, "GTM Push \u2014 Showcase & Sales Activation", "Go-To-Market")

add_task_table(doc, [
    ("**AI Capability Showcases** (bi-weekly) \u2014 all AI Power Team members should attend", "All AI Power Team"),
    ("Run **Sales Activations** targeting specific accounts", "Sales + Account Partner"),
    ("Assign **follow-up actions** with owners and due dates", "AI Champion"),
    ("**Assemble GTM Folder**: demo video, demo script, sales one-slider & sales deck", "AI Enablement Team"),
    ("**Publish to Capco MarketHub & AI Lab** to reach a larger audience", "AI Enablement Team"),
], 7)

doc.add_page_break()

# ============================================================
# STEP 8
# ============================================================
add_step_heading(doc, 8, "Track Results & Continuous Improvement", "Measure")

add_task_table(doc, [
    ("Track **CTAR** (Client Target Account Revenue): pipeline generated, deals won, revenue attributed to each AI solution", "AI Enablement + AI Champion + Sponsor"),
    ("Monitor **revenue metrics per solution**: total pipeline value, win rate, average deal size & time-to-close", "AI Enablement + Business Sponsor"),
    ("Monitor maturity: Idea \u2192 In Build \u2192 Demo-Ready \u2192 GTM-Ready \u2192 **In Pursuit**", "AI Enablement Team"),
    ("Capture **client feedback** & feed learnings back into next intake cycle", "AI Champion + Domain Leader"),
], 8)

# Continuous loop callout
table = doc.add_table(rows=1, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = table.rows[0].cells[0]
set_cell_shading(cell, LIGHT_COLORS[8])
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("\u21BB  Continuous Loop: ")
run.bold = True
run.font.size = Pt(10.5)
run.font.color.rgb = COLORS[8]
run = p.add_run("Results feed back into Step 1. Client feedback generates new ideas, successful demos inspire adjacent solutions, market learnings sharpen future intake.")
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# ============================================================
# SAVE
# ============================================================
output_path = "/home/user/Intake-to-GTM/AI-Use-Case-to-GTM-Step-Guide.docx"
doc.save(output_path)
print(f"DOCX saved to: {output_path}")
